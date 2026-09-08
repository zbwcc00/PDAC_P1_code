#!/usr/bin/env python3
"""Synchronize the published Zenodo DOI into the final Word manuscript."""

from __future__ import annotations

from pathlib import Path

from docx import Document


MANUSCRIPT = Path(r"D:\第二篇大论文\manuscript\PDAC_P1_full_manuscript_v11_audited_2026-09-07.docx")
REPLACEMENTS = {
    "All data analyzed are publicly available through GEO, TCGA/GDC, CPTAC, HPA, DepMap, PRISM, GTEx, the eQTL Catalogue, GWAS Catalog, and the cited resource portals[44,42,43]. Accession numbers and cohort roles are provided in Methods and Supplementary Table S1.": "All data analyzed are publicly available through GEO, TCGA/GDC, CPTAC, HPA, DepMap, PRISM, GTEx, the eQTL Catalogue, GWAS Catalog, and the cited resource portals[44,42,43]. Accession numbers and cohort roles are provided in Methods and Supplementary Table S1. Processed, non-identifying source-data tables supporting the main and supplementary figures and reported statistics are archived with the versioned code release in Zenodo (version v0.1.2, https://doi.org/10.5281/zenodo.22647158; concept DOI, https://doi.org/10.5281/zenodo.22647157). Raw third-party datasets are not redistributed in this archive.",
    "Analysis scripts, source-data tables, and manuscript-generation materials are retained in the project archive. A public archival repository and DOI should be supplied before submission.": "Analysis scripts, source-data tables, and manuscript-generation materials are available from Zenodo (version v0.1.2, https://doi.org/10.5281/zenodo.22647158) and the matching GitHub release (https://github.com/zbwcc00/PDAC_P1_code/releases/tag/v0.1.2).",
}


def main() -> None:
    document = Document(MANUSCRIPT)
    remaining = set(REPLACEMENTS)
    for paragraph in document.paragraphs:
        original = paragraph.text
        if original in REPLACEMENTS:
            paragraph.clear()
            paragraph.add_run(REPLACEMENTS[original])
            remaining.remove(original)
    if remaining:
        raise RuntimeError("Could not locate manuscript availability paragraph(s).")
    document.save(MANUSCRIPT)
    print(f"Updated {MANUSCRIPT}")


if __name__ == "__main__":
    main()
